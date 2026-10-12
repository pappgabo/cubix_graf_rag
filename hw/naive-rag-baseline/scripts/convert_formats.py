"""
Converts the plain-text F1 corpus into a MIX of real file formats
(pdf / docx / html / xlsx / txt) so the Naive RAG pipeline has to deal
with the document-conversion challenge (bonus task A).

- Each source .txt (with its metadata header) is rendered into a target
  format according to FORMAT_MAP, preserving the source metadata.
- A genuinely tabular XLSX ("ref-f1-champions.xlsx") is generated from
  curated F1 data (season results + driver summary).
- corpus/metadata.json is rebuilt to point at the new files and record
  each document's format.

Run AFTER download_corpus.py.
"""

import json
import html
from pathlib import Path

from docx import Document
from docx.shared import Pt
from fpdf import FPDF
from fpdf.enums import XPos, YPos
from openpyxl import Workbook

MC = dict(new_x=XPos.LMARGIN, new_y=YPos.NEXT)  # multi_cell: wrap to next line

CORPUS = Path(__file__).resolve().parent.parent / "corpus"

# Which target format each doc_id becomes. A few stay as .txt on purpose.
FORMAT_MAP = {
    # drivers
    "driver-lewis-hamilton": "pdf",
    "driver-max-verstappen": "pdf",
    "driver-michael-schumacher": "pdf",
    "driver-fernando-alonso": "docx",
    "driver-sebastian-vettel": "docx",
    "driver-nico-rosberg": "html",
    "driver-charles-leclerc": "html",
    "driver-kimi-raikkonen": "txt",
    "driver-jenson-button": "txt",
    "driver-valtteri-bottas": "txt",
    # teams
    "team-scuderia-ferrari": "pdf",
    "team-mclaren": "docx",
    "team-red-bull-racing": "docx",
    "team-mercedes-benz-in-formula-one": "html",
    "team-williams-racing": "html",
    "team-renault-in-formula-one": "txt",
    # seasons
    "season-2021-formula-one-world-championship": "pdf",
    "season-2008-formula-one-world-championship": "pdf",
    "season-2016-formula-one-world-championship": "docx",
    "season-2020-formula-one-world-championship": "html",
    "season-2012-formula-one-world-championship": "txt",
    # reference
    "reference-formula-one": "html",
    # the champions list becomes a real spreadsheet instead
    "reference-list-of-formula-one-world-drivers-champions": "xlsx",
}

# Windows Unicode font for PDF (handles en-dashes, accents, etc.)
FONT_REG = Path("C:/Windows/Fonts/arial.ttf")
FONT_BOLD = Path("C:/Windows/Fonts/arialbd.ttf")


def parse_doc(path: Path):
    """Return (meta_dict, body) from a header-stamped .txt file."""
    text = path.read_text(encoding="utf-8")
    head, sep, body = text.partition("=" * 70)
    meta = {}
    lines = head.splitlines()
    meta["title"] = lines[0].lstrip("# ").strip() if lines else path.stem
    for line in lines[1:]:
        if ": " in line and not line.startswith("#"):
            k, _, v = line.partition(": ")
            meta[k.strip()] = v.strip()
    return meta, body.strip()


def meta_lines(meta):
    """Human-readable source-metadata block reused in every format."""
    keys = ["doc_id", "category", "source", "url",
            "revision_id", "revision_timestamp", "retrieved_utc", "license"]
    return [f"{k}: {meta[k]}" for k in keys if k in meta]


def is_heading(line: str) -> bool:
    s = line.strip()
    return s.startswith("==") and s.endswith("==")


def heading_text(line: str) -> str:
    return line.strip().strip("=").strip()


# --- Writers ---------------------------------------------------------------

def write_pdf(meta, body, out: Path):
    pdf = FPDF()
    pdf.add_font("arial", "", str(FONT_REG))
    pdf.add_font("arial", "B", str(FONT_BOLD))
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_font("arial", "B", 16)
    pdf.multi_cell(0, 9, meta["title"], **MC)
    pdf.ln(1)

    pdf.set_font("arial", "", 8)
    pdf.set_text_color(110, 110, 110)
    for line in meta_lines(meta):
        pdf.multi_cell(0, 4, line, **MC)
    pdf.set_text_color(0, 0, 0)
    pdf.ln(3)

    for para in body.split("\n"):
        para = para.rstrip()
        if not para:
            pdf.ln(2)
            continue
        if is_heading(para):
            pdf.set_font("arial", "B", 12)
            pdf.ln(2)
            pdf.multi_cell(0, 6, heading_text(para), **MC)
            pdf.set_font("arial", "", 10)
        else:
            pdf.set_font("arial", "", 10)
            pdf.multi_cell(0, 5, para, **MC)
    pdf.output(str(out))


def write_docx(meta, body, out: Path):
    doc = Document()
    doc.add_heading(meta["title"], level=0)

    p = doc.add_paragraph()
    run = p.add_run("\n".join(meta_lines(meta)))
    run.font.size = Pt(8)
    run.italic = True

    for para in body.split("\n"):
        para = para.rstrip()
        if not para:
            continue
        if is_heading(para):
            doc.add_heading(heading_text(para), level=2)
        else:
            doc.add_paragraph(para)
    doc.save(str(out))


def write_html(meta, body, out: Path):
    parts = [
        "<!DOCTYPE html>",
        '<html lang="en"><head><meta charset="utf-8">',
        f"<title>{html.escape(meta['title'])}</title>",
    ]
    # keep source metadata machine-readable in <meta> tags too
    for k in ["doc_id", "category", "source", "url", "revision_id",
              "revision_timestamp", "retrieved_utc", "license"]:
        if k in meta:
            parts.append(f'<meta name="{k}" content="{html.escape(meta[k])}">')
    parts.append("</head><body>")
    parts.append(f"<h1>{html.escape(meta['title'])}</h1>")
    parts.append('<section class="source-metadata"><ul>')
    for line in meta_lines(meta):
        parts.append(f"<li>{html.escape(line)}</li>")
    parts.append("</ul></section>")

    for para in body.split("\n"):
        para = para.rstrip()
        if not para:
            continue
        if is_heading(para):
            parts.append(f"<h2>{html.escape(heading_text(para))}</h2>")
        else:
            parts.append(f"<p>{html.escape(para)}</p>")
    parts.append("</body></html>")
    out.write_text("\n".join(parts), encoding="utf-8")


def write_champions_xlsx(meta, out: Path):
    """A genuinely tabular spreadsheet built from curated F1 data."""
    wb = Workbook()

    # Sheet 1: results of the five seasons present in the corpus
    ws = wb.active
    ws.title = "Champions by season"
    ws.append(["Season", "World Champion", "Constructor", "Points"])
    for row in [
        [2008, "Lewis Hamilton", "McLaren", 98],
        [2012, "Sebastian Vettel", "Red Bull Racing", 281],
        [2016, "Nico Rosberg", "Mercedes", 385],
        [2020, "Lewis Hamilton", "Mercedes", 347],
        [2021, "Max Verstappen", "Red Bull Racing", 395.5],
    ]:
        ws.append(row)

    # Sheet 2: summary of the drivers covered in the corpus
    ws2 = wb.create_sheet("Drivers in corpus")
    ws2.append(["Driver", "Nationality", "World titles (as of 2024)",
                "Main constructors"])
    for row in [
        ["Lewis Hamilton", "British", 7, "McLaren; Mercedes"],
        ["Max Verstappen", "Dutch", 4, "Red Bull Racing"],
        ["Sebastian Vettel", "German", 4, "Red Bull Racing; Ferrari"],
        ["Fernando Alonso", "Spanish", 2, "Renault; McLaren"],
        ["Michael Schumacher", "German", 7, "Benetton; Ferrari"],
        ["Kimi Räikkönen", "Finnish", 1, "McLaren; Ferrari"],
        ["Nico Rosberg", "German", 1, "Williams; Mercedes"],
        ["Jenson Button", "British", 1, "BAR/Honda; Brawn; McLaren"],
        ["Charles Leclerc", "Monégasque", 0, "Sauber; Ferrari"],
        ["Valtteri Bottas", "Finnish", 0, "Williams; Mercedes"],
    ]:
        ws2.append(row)

    # Sheet 3: provenance so the source metadata survives the conversion
    ws3 = wb.create_sheet("Source metadata")
    ws3.append(["key", "value"])
    for k in ["title", "doc_id", "category", "source", "url",
              "revision_id", "revision_timestamp", "retrieved_utc", "license"]:
        if k in meta:
            ws3.append([k, meta[k]])
    ws3.append(["note", "Curated tabular data derived from the Wikipedia "
                        "List of Formula One World Drivers' Champions."])

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
    wb.save(str(out))


# --- Main ------------------------------------------------------------------

WRITERS = {"pdf": write_pdf, "docx": write_docx, "html": write_html}
EXT = {"pdf": ".pdf", "docx": ".docx", "html": ".html",
       "xlsx": ".xlsx", "txt": ".txt"}


def main():
    assert FONT_REG.exists(), f"Font not found: {FONT_REG}"
    metadata = []

    for txt in sorted(CORPUS.glob("*.txt")):
        meta, body = parse_doc(txt)
        doc_id = meta.get("doc_id", txt.stem)
        fmt = FORMAT_MAP.get(doc_id, "txt")
        out = CORPUS / f"{doc_id}{EXT[fmt]}"

        if fmt == "txt":
            pass  # leave the original .txt untouched
        elif fmt == "xlsx":
            write_champions_xlsx(meta, out)
            txt.unlink()
        else:
            WRITERS[fmt](meta, body, out)
            txt.unlink()

        print(f"  [{fmt:4s}] {out.name}")
        entry = {k: meta[k] for k in [
            "doc_id", "category", "source", "url", "revision_id",
            "revision_timestamp", "retrieved_utc", "license"] if k in meta}
        entry = {"file": out.name, "title": meta["title"],
                 "format": fmt, **entry}
        metadata.append(entry)

    (CORPUS / "metadata.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")

    counts = {}
    for m in metadata:
        counts[m["format"]] = counts.get(m["format"], 0) + 1
    print(f"\nDone. {len(metadata)} documents.")
    print("Formats:", ", ".join(f"{k}={v}" for k, v in sorted(counts.items())))


if __name__ == "__main__":
    main()
