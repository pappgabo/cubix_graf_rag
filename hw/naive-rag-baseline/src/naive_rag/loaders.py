"""Multi-format corpus loading (bonus task A).

Reads every document in the corpus directory (pdf / docx / html / xlsx / txt)
into LangChain `Document` objects while preserving source metadata
(file, doc_id, url, category, and page or section).

Heading-aware formats are converted to Markdown so section titles survive:
  docx -> pandoc (gfm)      html -> markdownify      txt -> '== X ==' -> '## X'
`ingest.py` then splits those documents by heading, so each chunk knows its
`section`. PDF keeps its page number; XLSX keeps the sheet name.
"""

from __future__ import annotations

import json
import re
import subprocess
from pathlib import Path

import openpyxl
from bs4 import BeautifulSoup
from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document
from markdownify import markdownify

from .config import get_settings
from .dependencies import require_for_formats, require_pandoc

# Formats converted to Markdown (heading-aware -> `section` metadata).
MARKDOWN_FORMATS = {"docx", "html", "txt"}

# metadata.json keys copied into every chunk. The manifest's own `source` field
# (publisher name) is deliberately not copied: `source` in the chunk metadata is
# the citation handle (URL or file name), set in `_base_meta`.
_META_KEYS = [
    "doc_id", "title", "category", "format", "url",
    "revision_id", "revision_timestamp", "license",
]
_HEADER_SEP = "=" * 70
# Provenance lines that conversion scripts stamped into the documents.
# They are stored in metadata.json / chunk metadata, so drop them from text.
_PROVENANCE = re.compile(
    r"^\W*(doc_id|category|source|url|revision_id|revision_timestamp|"
    r"retrieved_utc|license):")


def _base_meta(entry: dict) -> dict:
    meta = {k: entry[k] for k in _META_KEYS if k in entry}
    meta["file"] = entry["file"]
    meta["source"] = entry.get("url") or entry["file"]  # citation handle
    return meta


def _strip_provenance(text: str, scan_lines: int = 25) -> str:
    """Remove stamped provenance lines from the top of a document."""
    lines = text.splitlines()
    head = [ln for ln in lines[:scan_lines] if not _PROVENANCE.match(ln)]
    return "\n".join(head + lines[scan_lines:]).strip()


# --- per-format loaders ----------------------------------------------------

def _load_txt(path: Path, base: dict) -> list[Document]:
    text = path.read_text(encoding="utf-8")
    if _HEADER_SEP in text:
        text = text.split(_HEADER_SEP, 1)[1]
    # '== Section ==' -> '## Section', '=== Sub ===' -> '### Sub'
    text = re.sub(
        r"^(=+)\s*(.+?)\s*=+\s*$",
        lambda m: "#" * len(m.group(1)) + " " + m.group(2),
        text, flags=re.M)
    return [Document(page_content=text.strip(), metadata=dict(base))]


def _load_docx(path: Path, base: dict) -> list[Document]:
    out = subprocess.run(
        [require_pandoc(), "-f", "docx", "-t", "gfm", "--wrap=none", str(path)],
        capture_output=True, text=True, encoding="utf-8", check=True)
    text = _strip_provenance(out.stdout)
    return [Document(page_content=text, metadata=dict(base))]


def _load_html(path: Path, base: dict) -> list[Document]:
    soup = BeautifulSoup(path.read_text(encoding="utf-8"), "lxml")
    for node in soup.select(".source-metadata, script, style"):
        node.decompose()
    body = soup.body or soup
    text = markdownify(str(body), heading_style="ATX")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return [Document(page_content=_strip_provenance(text),
                     metadata=dict(base))]


def _load_pdf(path: Path, base: dict) -> list[Document]:
    out = []
    for d in PyPDFLoader(str(path)).load():
        page = d.metadata.get("page")
        text = d.page_content
        if page == 0:
            text = _strip_provenance(text)
        meta = {**base, "page": (page + 1) if isinstance(page, int) else page}
        out.append(Document(page_content=text, metadata=meta))
    return out


def _load_xlsx(path: Path, base: dict) -> list[Document]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        lines = [f"Sheet: {ws.title}"]
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
        if len(lines) > 1:
            meta = {**base, "section": f"sheet {ws.title}"}
            out.append(Document(page_content="\n".join(lines), metadata=meta))
    wb.close()
    return out


_LOADERS = {
    "pdf": _load_pdf, "docx": _load_docx, "html": _load_html,
    "xlsx": _load_xlsx, "txt": _load_txt,
}


def _format_of(entry: dict) -> str:
    return entry.get("format") or Path(entry["file"]).suffix.lstrip(".")


def load_corpus() -> list[Document]:
    """Load the whole corpus into LangChain Documents."""
    corpus_dir = get_settings().corpus_dir
    meta_path = corpus_dir / "metadata.json"
    if not meta_path.exists():
        raise FileNotFoundError(f"metadata.json not found in {corpus_dir}.")

    entries = json.loads(meta_path.read_text(encoding="utf-8"))
    # Fail fast: check external programs before parsing a single document.
    require_for_formats(_format_of(e) for e in entries
                        if (corpus_dir / e["file"]).exists())

    documents: list[Document] = []
    for entry in entries:
        path = corpus_dir / entry["file"]
        if not path.exists():
            print(f"  [WARN] missing file: {path.name}")
            continue
        loader = _LOADERS.get(_format_of(entry), _load_txt)
        documents.extend(loader(path, _base_meta(entry)))

    print(f"Loaded {len(documents)} source documents from {corpus_dir.name}.")
    return documents
